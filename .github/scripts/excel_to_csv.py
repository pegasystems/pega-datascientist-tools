#!/usr/bin/env python3
from __future__ import annotations

import argparse
from pathlib import Path

import openpyxl
from openpyxl.utils.cell import range_boundaries
import polars as pl


def main():
    p = argparse.ArgumentParser()
    p.add_argument(
        "--input",
        required=True,
        help="Path to .xlsx",
    )
    p.add_argument(
        "--table",
        required=True,
        help="Named Excel table (e.g. MyTable)",
    )
    p.add_argument(
        "--output",
        required=True,
        help="Path to output CSV",
    )
    args = p.parse_args()

    wb = openpyxl.load_workbook(args.input, data_only=True)

    # Find the named table
    for ws in wb.worksheets:
        if args.table in ws.tables:
            ref = ws.tables[args.table].ref
            break
    else:
        raise ValueError(f"Table '{args.table}' not found in {args.input}")

    # Extract table data
    min_col, min_row, max_col, max_row = range_boundaries(ref)
    rows = [
        list(r)
        for r in ws.iter_rows(
            min_row=min_row,
            max_row=max_row,
            min_col=min_col,
            max_col=max_col,
            values_only=True,
        )
    ]

    if len(rows) < 2:
        raise ValueError(f"Table '{args.table}' has no data (range {ws.title}!{ref}).")

    # Build Polars frame and write CSV
    pl_df = pl.DataFrame(rows[1:], schema=[str(h) for h in rows[0]], orient="row")
    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    pl_df.write_csv(args.output, line_terminator="\n")


if __name__ == "__main__":
    main()
